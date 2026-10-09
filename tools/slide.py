#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SLIDE: dựng bài trình chiếu (file .json) ra PPTX sửa được trên Google Slides, PowerPoint, Keynote, kèm kiểm và xem trước.

    python3 tools/slide.py <dự án>/slide/<tên>.json --nhap   # nháp: PPTX + ảnh xem trước từng slide + tờ tổng thể + báo cáo
    python3 tools/slide.py <dự án>/slide/<tên>.json          # bản cuối: PPTX + PDF chiếu dự phòng vào Thanh pham/<dự án>/NN <tên>/
    python3 tools/slide.py --kiem                            # cổng kiểm lõi trình chiếu: dựng mọi deck mẫu trinh-chieu/mau/*.json
    python3 tools/slide.py --xuat-phong <thư mục>            # xuất phông TTF đủ dấu để cài trên máy dùng PowerPoint, Keynote

Một nguồn với lõi ấn phẩm: màu và phông theo brand/brand.json (chủ đề sáng cho nội dung, chủ đề tối cùng thương hiệu
cho bìa, chuyển phần, kết), sơ đồ cùng dữ liệu với minh-hoa/so-do.js (trinh-chieu/so-do-pptx.js vẽ bằng hình khối gốc),
biểu tượng từ he-thong/bieu-tuong.css, luật chữ từ tools/ve.py và phong-cach/tu-ngu.json, hình đặc thù dựng bằng lõi HTML
(trường "anPham": một file thiet-ke/*.json, ve.py dựng ra PNG rồi chèn vào slide).

Máy kiểm (báo cáo <tên>-bao-cao.json, chuan/09 mục 6):
  LỖI: tran-chu (chữ không vừa hộp ở cỡ tối thiểu), gian-chu (thuộc tính giãn chữ spc: vỡ chữ trên Google Slides),
       so-do, loi-js, kieu, anh (thiếu ảnh), tu-ngu, gach-dai
  XEM: chu-nho, nhieu-chu, thieu-hinh, tuong-phan, thieu-nguon, phong-la, goi-y-chu, title-case
Chạy ở sandbox đám mây (Node có pptxgenjs, fontkit, sharp; LibreOffice để xem trước). Thiếu thư viện Node thì tự cài vào
vùng tạm NGOÀI repo (<thuMucDuAn>/_tam/node).
"""
import argparse
import copy
import datetime
import glob
import json
import os
import re
import shutil
import subprocess
import sys
sys.dont_write_bytecode = True  # không để __pycache__ trong repo
import tempfile
import zipfile

TOOLS = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.dirname(TOOLS)
sys.path.insert(0, TOOLS)
import ve  # noqa: E402

TC = os.path.join(GOC, 'trinh-chieu')
LOI = {'tran-chu', 'gian-chu', 'xml', 'so-do', 'loi-js', 'kieu', 'anh', 'tu-ngu', 'gach-dai'}
THU_VIEN = {'pptxgenjs': 'pptxgenjs@4.0.1', 'fontkit': 'fontkit@2', 'sharp': 'sharp'}
PHONG_HOP_LE = {'Be Vietnam Pro', 'Lora', 'Playfair Display', 'JetBrains Mono'}
KHUNG_IN = (13.333, 7.5)


# ---------------------------------------------------------------- môi trường Node
def node_path():
    ung = [os.environ.get('NODE_PATH', ''), ve.thu_muc_tam('node', 'node_modules'), '/opt/npm-tools/node_modules']
    try:
        ung.append(subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True, timeout=30).stdout.strip())
    except Exception:
        pass
    ds = [p for p in ':'.join(ung).split(':') if p and os.path.isdir(p)]
    thieu = [k for k in THU_VIEN if not any(os.path.isdir(os.path.join(p, k)) for p in ds)]
    if thieu:
        noi = ve.thu_muc_tam('node')
        ve.chan_ghi_repo(noi, 'thư viện Node')
        os.makedirs(noi, exist_ok=True)
        if not os.path.exists(os.path.join(noi, 'package.json')):
            with open(os.path.join(noi, 'package.json'), 'w') as f:
                f.write('{"name":"tam-trinh-chieu","private":true}')
        print(f'  cài thư viện Node vào vùng tạm: {", ".join(thieu)}')
        subprocess.run(['npm', 'i', '--silent', '--no-audit', '--no-fund'] + [THU_VIEN[k] for k in thieu], cwd=noi, check=True)
        ds.insert(0, os.path.join(noi, 'node_modules'))
    return ':'.join(dict.fromkeys(ds))


# ---------------------------------------------------------------- phông cho bản xem trước (LibreOffice)
def cai_phong_xem_truoc(dich=None):
    """Gộp các tệp con woff2 (latin, latin-ext, vietnamese) của fonts/ thành TTF đủ dấu: cài cho LibreOffice ở sandbox,
    hoặc xuất ra `dich` để người dùng cài trên máy (PowerPoint, Keynote không đọc woff2)."""
    cai = dich is None
    dich = dich or os.path.expanduser('~/.local/share/fonts/xuong-thiet-ke')
    dau = os.path.join(dich, '.xong')
    if cai and os.path.exists(dau):
        return
    try:
        from fontTools.ttLib import TTFont
        from fontTools.merge import Merger
    except Exception:
        print('  (thiếu fontTools: bản xem trước sẽ dùng phông thay thế)')
        return
    os.makedirs(dich, exist_ok=True)
    tam = tempfile.mkdtemp(prefix='phong-')
    for ho in ('be-vietnam-pro', 'lora', 'playfair-display'):
        nhom = {}
        for p in glob.glob(os.path.join(GOC, 'fonts', ho, '*.woff2')):
            m = re.match(rf'{ho}-(latin|latin-ext|vietnamese)-(\d+)-(normal|italic)\.woff2$', os.path.basename(p))
            if m:
                nhom.setdefault((m.group(2), m.group(3)), []).append(p)
        for (w, st), ds in nhom.items():
            ttf = []
            for p in sorted(ds, key=lambda x: ('vietnamese' not in x, 'ext' in x)):
                try:
                    t = TTFont(p)
                    t.flavor = None
                    q = os.path.join(tam, os.path.basename(p) + '.ttf')
                    t.save(q)
                    ttf.append(q)
                except Exception:
                    pass
            if not ttf:
                continue
            try:
                Merger().merge(ttf).save(os.path.join(dich, f'{ho}-{w}-{st}.ttf'))
            except Exception:
                shutil.copy(ttf[0], os.path.join(dich, f'{ho}-{w}-{st}.ttf'))
    shutil.rmtree(tam, ignore_errors=True)
    if cai:
        subprocess.run(['fc-cache', '-f', dich], capture_output=True)
        open(dau, 'w').close()
    else:
        print(f'Phông TTF đã xuất: {dich} (bấm đúp từng tệp để cài; giấy phép SIL OFL ở fonts/*/OFL.txt)')


# ---------------------------------------------------------------- chuẩn bị deck
SO_MUC = {'chuoi': 'buoc', 'vong-lap': 'buoc', 'hanh-trinh': 'tram', 'tang': 'tang', 'trung-tam': 'nhanh', 'the-luoi': 'the',
          'radar': 'truc', 'venn': 'tap', 'dong-tam': 'lop', 'pho': 'doan', 'xoan-oc': 'buoc', 'isotype': 'hang'}


def so_muc(sd):
    if sd.get('loai') in SO_MUC:
        return len(sd.get(SO_MUC[sd['loai']]) or [])
    return {'ma-tran': 4, 'so-sanh': 2}.get(sd.get('loai'), 0)


def trai_trinh_tu(ds):
    """Slide sơ đồ có "trinhTu": true sinh thêm mỗi mục một slide báo hiệu mục đó (chuan/06 R6.3; Mayer: báo hiệu, phân đoạn)."""
    ra = []
    for s in ds:
        sd = s.get('soDo') or ((s.get('hinh') or {}).get('soDo'))
        if not s.get('trinhTu') or not sd:
            ra.append(s)
            continue
        if s.get('trinhTu') != 'chi-tung-muc':
            ra.append({k: v for k, v in s.items() if k != 'trinhTu'})
        tieu = s.get('tieuDeTungMuc') or []
        ghi = s.get('ghiChuTungMuc') or []
        for i in range(so_muc(sd)):
            m = copy.deepcopy({k: v for k, v in s.items() if k not in ('trinhTu', 'tieuDeTungMuc', 'ghiChuTungMuc')})
            (m.get('soDo') or m['hinh']['soDo'])['mucNhan'] = i
            if i < len(tieu) and tieu[i]:
                m['tieuDe'] = tieu[i]
            m['ghiChu'] = ghi[i] if i < len(ghi) else ''
            ra.append(m)
    return ra


def duong_anh(src, du_an):
    if not src:
        return None
    if isinstance(src, dict):
        src = src.get('src')
    if src.startswith('kho:'):
        return os.path.join(GOC, src[4:])
    return src if os.path.isabs(src) else os.path.join(du_an, src)


def dung_an_pham(ap, du_an, w_in, h_in, tam):
    """Hình đặc thù dựng bằng lõi HTML (ve.py) ra PNG đúng tỉ lệ ô trên slide."""
    spec = duong_anh(ap, du_an)
    W, H = round(w_in * 160), round(h_in * 160)
    ra = os.path.join(tam, 'an-pham', os.path.splitext(os.path.basename(spec))[0] + f'-{W}x{H}')
    try:
        ve.ve_an_pham(spec, ds_tt=[f'tu-do:{W}x{H}'], ra=ra, nhap=False, tong_the=False, im_lang=True)
    except SystemExit as e:
        return None, str(e)
    except Exception as e:  # noqa
        return None, repr(e)
    ds = sorted(glob.glob(os.path.join(ra, '*.png')))
    return (ds[0] if ds else None), None


def chuan_bi(deck, du_an, tam, canh):
    ds = trai_trinh_tu(deck.get('slide') or [])
    for i, s in enumerate(ds):
        k = s.get('kieu', 'y-chinh')
        if s.get('anh') and k in ('bia', 'anh', 'hinh'):
            s['tepAnh'] = duong_anh(s['anh'], du_an)
            if isinstance(s['anh'], dict) and s['anh'].get('tieuDiem'):
                s['tieuDiem'] = s['anh']['tieuDiem']
        h = s.get('hinh')
        if isinstance(h, dict) and h.get('anh'):
            h['tepAnh'] = duong_anh(h['anh'], du_an)
        for chu_the, khung in ((s, (11.73, 5.2)), (h if isinstance(h, dict) else None, (6.1, 5.0))):
            if chu_the and chu_the.get('anPham'):
                png, loi = dung_an_pham(chu_the['anPham'], du_an, *khung, tam)
                if png:
                    chu_the['tepAnh'] = png
                else:
                    canh.append({'trang': i + 1, 'loai': 'anh', 'id': 'anPham', 'chiTiet': f'không dựng được {chu_the["anPham"]}: {loi}'})
    deck = dict(deck, slide=ds)
    return deck


def chu_de_cap(deck):
    th, sang = ve.tron_chu_de(deck)
    b = ve.brand()
    if deck.get('chuDeToi'):
        toi = ve.tron_chu_de(dict(deck, chuDe=deck['chuDeToi']))[1]
    elif sang.get('toi'):
        toi = sang
    else:
        ung = [c for c in (th.get('chuDeHop') or []) if b['chuDe'].get(c, {}).get('toi')]
        toi = ve.tron_chu_de(dict(deck, chuDe=ung[0]))[1] if ung else sang
    return th, sang, toi


# ---------------------------------------------------------------- sửa và kiểm sau khi dựng
_PPR = re.compile(r'<a:pPr\b[^>]*?(?:/>|>.*?</a:pPr>)', re.S)


def _mot_ppr(m):
    ds = list(_PPR.finditer(m.group(1)))
    if len(ds) < 2:
        return m.group(0)
    giu = ds[0].group(0)
    than = _PPR.sub('', m.group(1))
    return '<a:p>' + giu + than + '</a:p>'


def sua_pptx(pptx):
    """PptxGenJS ghi một <a:pPr> cho MỖI đoạn chạy chữ trong cùng đoạn văn (đoạn có chữ đậm, nghiêng nhấn): sai lược đồ
    OOXML, PowerPoint có thể đòi sửa tệp, Google Slides và LibreOffice lấy định dạng đoạn cuối (mất gạch đầu dòng).
    Giữ <a:pPr> đầu tiên của mỗi đoạn văn."""
    tam = pptx + '.tam'
    doi = 0
    with zipfile.ZipFile(pptx) as zi, zipfile.ZipFile(tam, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            b = zi.read(it.filename)
            if re.match(r'ppt/(slides/slide|notesSlides/notesSlide)\d+\.xml$', it.filename):
                x = b.decode('utf-8')
                y = re.sub(r'<a:p>(.*?)</a:p>', _mot_ppr, x, flags=re.S)
                if y != x:
                    doi += 1
                    b = y.encode('utf-8')
            zo.writestr(it, b)
    os.replace(tam, pptx)
    return doi


def kiem_pptx(pptx):
    ra = []
    with zipfile.ZipFile(pptx) as z:
        xml = ''.join(z.read(n).decode('utf-8', 'ignore') for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n))
    gian = [v for v in re.findall(r'\bspc="(-?\d+)"', xml) if v != '0']
    if gian:
        ra.append({'trang': 0, 'loai': 'gian-chu', 'id': 'spc', 'chiTiet': f'{len(gian)} chỗ giãn chữ (spc): Google Slides sẽ chồng chữ tiếng Việt; bỏ mọi charSpacing'})
    thua = sum(1 for p in re.findall(r'<a:p>(.*?)</a:p>', xml, flags=re.S) if len(_PPR.findall(p)) > 1)
    if thua:
        ra.append({'trang': 0, 'loai': 'xml', 'id': 'pPr', 'chiTiet': f'{thua} đoạn văn có nhiều định dạng đoạn (pPr): tệp sai lược đồ'})
    phong = set(re.findall(r'<a:latin typeface="([^"]+)"', xml)) - PHONG_HOP_LE
    if phong:
        ra.append({'trang': 0, 'loai': 'phong-la', 'id': 'phong', 'chiTiet': f'phông ngoài hệ của xưởng: {", ".join(sorted(phong))}'})
    return ra


def xem_truoc(pptx, ra_dir, ten):
    """PPTX -> PDF (LibreOffice) -> JPG từng slide. Trả (pdf, [jpg])."""
    if not shutil.which('soffice'):
        print('  (không có LibreOffice: bỏ bước xem trước)')
        return None, []
    cai_phong_xem_truoc()
    tam = tempfile.mkdtemp(prefix='xem-')
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tam, pptx], capture_output=True, timeout=300)
    pdf = os.path.join(tam, os.path.splitext(os.path.basename(pptx))[0] + '.pdf')
    if not os.path.exists(pdf):
        print('  (LibreOffice không chuyển được sang PDF)')
        return None, []
    xem = os.path.join(ra_dir, 'xem')
    os.makedirs(xem, exist_ok=True)
    for p in glob.glob(os.path.join(xem, '*.jpg')):
        os.remove(p)
    subprocess.run(['pdftoppm', '-jpeg', '-r', '96', pdf, os.path.join(xem, ten)], capture_output=True, timeout=300)
    return pdf, sorted(glob.glob(os.path.join(xem, '*.jpg')))


def doc_pptx(pptx):
    """Chữ và ghi chú người nói từng slide của một PPTX bất kỳ (python-pptx). Trả [{trang, chu: [], ghiChu}]."""
    try:
        from pptx import Presentation
    except Exception:
        return []
    ra = []
    for i, sl in enumerate(Presentation(pptx).slides, 1):
        chu = []
        for sh in sl.shapes:
            khung = [sh] if sh.has_text_frame else []
            if getattr(sh, 'has_table', False) and sh.has_table:
                khung = [c for r in sh.table.rows for c in r.cells]
            for k in khung:
                for para in k.text_frame.paragraphs:
                    t = ''.join(r.text for r in para.runs).strip()
                    if t:
                        chu.append(('  ' * para.level) + t)
        ghi = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ''
        ra.append({'trang': i, 'chu': chu, 'ghiChu': ghi})
    return ra


def xem_pptx(pptx, ra=None):
    """Bài làm ngoài xưởng: kiểm tương thích, ảnh xem trước, tờ tổng thể, chép chữ và ghi chú ra <tên>-chu.md."""
    ten = os.path.splitext(os.path.basename(pptx))[0]
    ra = ra or ve.thu_muc_tam('xem-pptx', ten)
    ve.chan_ghi_repo(ra, 'thư mục ra')
    os.makedirs(ra, exist_ok=True)
    canh = kiem_pptx(pptx)
    trang = doc_pptx(pptx)
    with open(os.path.join(ra, f'{ten}-chu.md'), 'w', encoding='utf-8') as f:
        for t in trang:
            f.write(f'## Slide {t["trang"]}\n\n' + '\n'.join(t['chu']) + '\n')
            if t['ghiChu']:
                f.write('\nGhi chú người nói: ' + t['ghiChu'] + '\n')
            f.write('\n')
    pdf, anh = xem_truoc(pptx, ra, ten)
    if anh:
        tt = [{'kieu': '', 'soTieng': ''} for _ in anh]
        to_tong_the(anh, tt, os.path.join(ra, f'{ten}-tong-the.jpg'), f'{ten} - {len(anh)} slide (bài có sẵn)', '')
    print(f'  {len(trang) or len(anh)} slide, {len(canh)} cảnh báo')
    for c in canh:
        print(f'      - [{"LỖI " if c["loai"] in LOI else ""}{c["loai"]}] {c["chiTiet"]}')
    print(f'Chữ và ghi chú: {os.path.join(ra, ten + "-chu.md")}' + (f'\nẢnh xem trước, tờ tổng thể: {ra}' if anh else ''))
    return canh


def to_tong_the(anh, trang, ra, tieu_de, don_vi='tiếng'):
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return None
    if not anh:
        return None
    cot, wT = 4, 640
    hT = round(wT * 7.5 / 13.333)
    le, khe, nhan = 40, 28, 34
    hang = (len(anh) + cot - 1) // cot
    W = le * 2 + cot * wT + (cot - 1) * khe
    H = 90 + hang * (hT + nhan + khe) + le
    to = Image.new('RGB', (W, H), (236, 233, 226))
    d = ImageDraw.Draw(to)
    try:
        f = ImageFont.truetype(ve._phong_he_thong(), 22)
    except Exception:
        f = ImageFont.load_default()
    d.text((le, 30), tieu_de, fill=(40, 40, 40), font=f)
    for i, p in enumerate(anh):
        im = Image.open(p).convert('RGB').resize((wT, hT), Image.LANCZOS)
        x = le + (i % cot) * (wT + khe)
        y = 90 + (i // cot) * (hT + nhan + khe)
        t = trang[i] if i < len(trang) else {}
        d.text((x, y + 4), f'{i + 1:02d}  {t.get("kieu", "")}  {t.get("soTieng", "")} {don_vi}', fill=(60, 60, 60), font=f)
        to.paste(im, (x, y + nhan))
    to.save(ra, 'JPEG', quality=86)
    return ra


# ---------------------------------------------------------------- chạy
def dung_slide(deck_path, ra=None, nhap=False, xem=True, im_lang=False):
    deck = ve.doc_json(deck_path)
    ten = os.path.splitext(os.path.basename(deck_path))[0]
    d = os.path.dirname(os.path.abspath(deck_path))
    du_an = os.path.dirname(d) if os.path.basename(d) in ('slide', 'thiet-ke') else d
    ra_mac_dinh = not ra
    ra = ra or ve.thu_muc_ra(du_an, ten, nhap)
    ve.chan_ghi_repo(ra, 'thư mục ra')
    ho_so = ve.thu_muc_ra(du_an, ten, True) if (ra_mac_dinh and not nhap) else ra
    ve.chan_ghi_repo(ho_so, 'thư mục hồ sơ (tờ tổng thể, báo cáo)')
    os.makedirs(ra, exist_ok=True)
    os.makedirs(ho_so, exist_ok=True)
    tam = ve.thu_muc_tam('trinh-chieu', ten)
    ve.chan_ghi_repo(tam, 'vùng tạm')
    os.makedirs(tam, exist_ok=True)

    canh = []
    ngon_ngu = deck.get('ngonNgu', 'vi')
    for c in ve.kiem_chu({'noiDung': {'slide': deck.get('slide') or [], 'chanTrang': deck.get('chanTrang') or {}}}):
        m = re.match(r'slide\.(\d+)\.(.*)', c['id'] or '')
        canh.append({'trang': int(m.group(1)) + 1 if m else 0, 'loai': 'gach-dai' if '—' in c['chiTiet'][:40] and c['loai'] == 'tu-ngu' else c['loai'],
                     'id': m.group(2) if m else c['id'], 'chiTiet': c['chiTiet'] + (' (số trang theo file nguồn)' if m else '')})
    canh = [c for c in canh if not (c['loai'] == 'title-case' and re.search(r'(^|\.)(ngay|nguoi|lienHe|tacGia|nguon|chu|kicker)(\.\d+)?$', c['id'] or ''))]
    deck2 = chuan_bi(deck, du_an, tam, canh)
    th, sang, toi = chu_de_cap(deck)
    pptx = os.path.join(ra, f'{ten}.pptx')
    bc_node = os.path.join(tam, 'bao-cao-node.json')
    goi = {'deck': deck2, 'ra': pptx, 'baoCao': bc_node, 'kho': GOC, 'fonts': os.path.join(GOC, 'fonts'), 'tam': tam,
           'ngonNgu': ngon_ngu, 'chuDe': {'sang': sang, 'toi': toi}, 'thuongHieu': th}
    p_goi = os.path.join(tam, 'goi.json')
    with open(p_goi, 'w', encoding='utf-8') as f:
        json.dump(goi, f, ensure_ascii=False)
    env = dict(os.environ, NODE_PATH=node_path())
    kq = subprocess.run(['node', os.path.join(TC, 'dung.js'), p_goi], env=env, capture_output=True, text=True, timeout=600)
    if kq.returncode != 0:
        raise SystemExit('Dựng PPTX lỗi:\n' + (kq.stderr or kq.stdout)[-3000:])
    bc = ve.doc_json(bc_node)
    sua_pptx(pptx)
    canh += bc['canhBao']
    canh += kiem_pptx(pptx)

    pdf, anh = (None, [])
    if xem or not nhap:
        pdf, anh = xem_truoc(pptx, ho_so, ten)
    bao_cao = {'deck': deck_path, 'luc': datetime.datetime.now().isoformat(timespec='seconds'), 'banNhap': bool(nhap),
               'pptx': pptx, 'soSlide': bc['soSlide'], 'chuDe': {'sang': sang['id'], 'toi': toi['id']},
               'trang': bc['trang'], 'canhBao': canh}
    if pdf and not nhap:
        dich = os.path.join(ra, f'{ten}.pdf')
        shutil.copy(pdf, dich)
        bao_cao['pdf'] = dich
    if anh:
        tt = to_tong_the(anh, bc['trang'], os.path.join(ho_so, f'{ten}-tong-the.jpg'), f'{ten} - {bc["soSlide"]} slide - {sang["id"]}/{toi["id"]} - {bao_cao["luc"]}', 'words' if ngon_ngu == 'en' else 'tiếng')
        if tt:
            bao_cao['tongThe'] = tt
        bao_cao['xem'] = anh
    with open(os.path.join(ho_so, f'{ten}-bao-cao.json'), 'w', encoding='utf-8') as f:
        json.dump(bao_cao, f, ensure_ascii=False, indent=1)
    loi = [c for c in canh if c['loai'] in LOI]
    if not im_lang:
        print(f'  {bc["soSlide"]} slide, chủ đề {sang["id"]} / {toi["id"]}, {len(canh)} cảnh báo ({len(loi)} LỖI)')
        for c in sorted(canh, key=lambda c: (c['trang'], c['loai'])):
            print(f'      - [{"LỖI " if c["loai"] in LOI else ""}{c["loai"]}] trang {c["trang"]} {c["id"]}: {c["chiTiet"]}')
        print(f'PPTX: {pptx}' + (f'\nPDF chiếu dự phòng: {bao_cao["pdf"]}' if bao_cao.get('pdf') else '')
              + (f'\nTờ tổng thể, báo cáo: {ho_so}' if anh else ''))
    return bao_cao


def kiem():
    ra_goc = ve.thu_muc_tam('kiem-slide')
    ve.chan_ghi_repo(ra_goc, 'thư mục kiểm')
    loi_tong = 0
    for p in sorted(glob.glob(os.path.join(TC, 'mau', '*.json'))):
        ten = os.path.splitext(os.path.basename(p))[0]
        print(f'== {ten}')
        bc = dung_slide(p, ra=os.path.join(ra_goc, ten), nhap=True, xem=True, im_lang=False)
        loi_tong += sum(1 for c in bc['canhBao'] if c['loai'] in LOI)
    print(f'\nKIỂM TRÌNH CHIẾU: {"ĐẠT" if loi_tong == 0 else f"CHƯA ĐẠT ({loi_tong} LỖI)"}  (tờ tổng thể ở {ra_goc}: nhìn trước khi dùng)')
    return loi_tong == 0


def main():
    ap = argparse.ArgumentParser(description='Dựng bài trình chiếu ra PPTX, kèm kiểm và xem trước.')
    ap.add_argument('deck', nargs='?', help='file trình chiếu .json; hoặc một .pptx có sẵn: chỉ kiểm, xem trước, chép chữ và ghi chú')
    ap.add_argument('--nhap', action='store_true', help='bản nháp vào <dự án>/nhap/<tên>/ kèm ảnh xem trước, tờ tổng thể')
    ap.add_argument('--ra', help='thư mục ra (mặc định: nháp <dự án>/nhap/<tên>/, bản cuối <thuMucThanhPham>/<dự án>/NN <tên>/)')
    ap.add_argument('--khong-xem', action='store_true', help='bỏ bước xem trước bằng LibreOffice')
    ap.add_argument('--kiem', action='store_true', help='cổng kiểm lõi: dựng mọi deck mẫu trong trinh-chieu/mau/')
    ap.add_argument('--xuat-phong', help='xuất phông TTF đủ dấu ra thư mục này (ngoài repo)')
    a = ap.parse_args()
    if a.xuat_phong:
        ve.chan_ghi_repo(a.xuat_phong, 'thư mục phông')
        os.makedirs(a.xuat_phong, exist_ok=True)
        cai_phong_xem_truoc(a.xuat_phong)
        return
    if a.kiem:
        sys.exit(0 if kiem() else 1)
    if not a.deck:
        ap.error('cần đường dẫn file trình chiếu (.json), một tệp .pptx có sẵn để xem, hoặc --kiem')
    if a.deck.lower().endswith('.pptx'):
        xem_pptx(a.deck, ra=a.ra)
        return
    dung_slide(a.deck, ra=a.ra, nhap=a.nhap, xem=not a.khong_xem)


if __name__ == '__main__':
    main()
